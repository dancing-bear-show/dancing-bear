"""Tests for SlideGenerator method delegation — testing via the unified SlideGenerator class.

These classes use names that conflict with the per-module mixin test files
(test__table.py, test__styling.py, test__shape_utils.py), which test the same
underlying methods via their mixin interfaces. These tests exercise the same
methods as exposed through SlideGenerator.
"""

import unittest
from unittest.mock import MagicMock, patch

from lxml import etree
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

from slides._styling import TextStyle
from slides.constants import (
    DEFAULT_SECTION_HEADER_MAX_LENGTH,
    FONT_SIZE_HEADER,
    HIGHLIGHT_THEME_COLOR,
    LINK_BLUE,
    SPACING_AFTER_BULLET,
    SPACING_AFTER_HEADER,
    SPACING_BEFORE_BULLET,
    SPACING_BEFORE_HEADER,
    TABLE_HEADER_BG,
    TABLE_ROW_EVEN_BG,
    TABLE_ROW_ODD_BG,
    VERTICAL_ANCHOR_MIDDLE,
)
from slides.generator import SlideGenerator
from slides.schema import BulletItem


class TestSlideGeneratorIsSectionHeader(unittest.TestCase):
    """Tests for SlideGenerator._is_section_header method."""

    def setUp(self):
        """Create a SlideGenerator instance for testing."""
        self.generator = SlideGenerator(template_path="/fake/template.pptx")

    def test_ends_with_colon_is_header(self):
        """Text ending with colon and under 40 chars is a header."""
        self.assertTrue(self.generator._is_section_header("Overview:"))
        self.assertTrue(self.generator._is_section_header("Key Findings:"))
        self.assertTrue(self.generator._is_section_header("Action Items:"))

    def test_no_colon_not_header(self):
        """Text without colon is not a header."""
        self.assertFalse(self.generator._is_section_header("Overview"))
        self.assertFalse(self.generator._is_section_header("This is a regular bullet point"))

    def test_long_text_with_colon_not_header(self):
        """Text over 40 chars with colon is not a header (sentence with colon)."""
        long_text = "This is a very long sentence that contains a colon here:"
        self.assertFalse(self.generator._is_section_header(long_text))
        self.assertGreater(len(long_text), DEFAULT_SECTION_HEADER_MAX_LENGTH)

    def test_empty_string_not_header(self):
        """Empty string is not a header."""
        self.assertFalse(self.generator._is_section_header(""))

    def test_whitespace_only_not_header(self):
        """Whitespace-only string is not a header."""
        self.assertFalse(self.generator._is_section_header("   "))
        self.assertFalse(self.generator._is_section_header("\t\n"))

    def test_colon_in_middle_not_header(self):
        """Colon in middle of text is not a header."""
        self.assertFalse(self.generator._is_section_header("Time: 10:30 AM"))

    def test_exactly_at_limit_with_colon_is_header(self):
        """Text exactly at boundary is still a header (< 40 check)."""
        # 38 chars + colon = 39 total, should pass < 40 check
        text = "A" * 38 + ":"  # 39 chars total
        self.assertTrue(self.generator._is_section_header(text))

    def test_at_limit_with_colon_not_header(self):
        """Text at 40 chars is not a header (must be < 40)."""
        text = "A" * 39 + ":"  # 40 chars total
        self.assertFalse(self.generator._is_section_header(text))

    def test_none_not_header(self):
        """None value returns False (if applicable)."""
        # The method should handle falsy values
        self.assertFalse(self.generator._is_section_header(None))


class TestSlideGeneratorFormatBulletText(unittest.TestCase):
    """Tests for SlideGenerator._format_bullet_text method.

    Note: _format_bullet_text strips existing bullet characters from text
    because native PowerPoint bullets are applied via XML, not text characters.
    """

    def setUp(self):
        """Create a SlideGenerator instance for testing."""
        self.generator = SlideGenerator(template_path="/fake/template.pptx")

    def test_plain_text_returned_as_is(self):
        """Plain text without bullet chars is returned cleaned (stripped)."""
        result = self.generator._format_bullet_text("Plain text", 0)
        self.assertEqual(result, "Plain text")

    def test_level_0_returns_clean_text(self):
        """Level 0 returns text without bullet characters."""
        result = self.generator._format_bullet_text("Level zero", 0)
        self.assertEqual(result, "Level zero")

    def test_level_1_returns_clean_text(self):
        """Level 1 returns text without bullet characters."""
        result = self.generator._format_bullet_text("Level one", 1)
        self.assertEqual(result, "Level one")

    def test_level_2_returns_clean_text(self):
        """Level 2 returns text without bullet characters."""
        result = self.generator._format_bullet_text("Level two", 2)
        self.assertEqual(result, "Level two")

    def test_level_beyond_max_returns_clean_text(self):
        """Levels beyond defined return clean text."""
        result = self.generator._format_bullet_text("Deep level", 5)
        self.assertEqual(result, "Deep level")

    def test_existing_bullet_chars_are_stripped(self):
        """Text with existing bullet characters gets them stripped."""
        test_cases = [
            ("• Already has bullet", "Already has bullet"),
            ("◦ Open circle bullet", "Open circle bullet"),
            ("▪ Square bullet", "Square bullet"),
            ("‣ Triangle bullet", "Triangle bullet"),
            ("- Dash bullet", "Dash bullet"),
            ("* Asterisk bullet", "Asterisk bullet"),
        ]
        for text, expected in test_cases:
            result = self.generator._format_bullet_text(text, 0)
            self.assertEqual(result, expected, f"Failed for: {text}")

    def test_section_header_returned_as_is(self):
        """Section headers (ending with :) are returned as-is."""
        result = self.generator._format_bullet_text("Overview:", 0)
        self.assertEqual(result, "Overview:")

    def test_empty_string_unchanged(self):
        """Empty string is returned as-is."""
        result = self.generator._format_bullet_text("", 0)
        self.assertEqual(result, "")

    def test_whitespace_only_unchanged(self):
        """Whitespace-only text is returned as-is."""
        result = self.generator._format_bullet_text("   ", 0)
        self.assertEqual(result, "   ")

    def test_text_with_leading_whitespace(self):
        """Text with leading whitespace gets trimmed."""
        result = self.generator._format_bullet_text("  Indented text  ", 0)
        self.assertEqual(result, "Indented text")


class TestSlideGeneratorGetThemeColor(unittest.TestCase):
    """Tests for SlideGenerator._get_theme_color method."""

    def setUp(self):
        """Create a SlideGenerator instance for testing."""
        self.generator = SlideGenerator(template_path="/fake/template.pptx")

    def test_valid_color_names(self):
        """Valid color names return corresponding MSO_THEME_COLOR."""
        self.assertEqual(self.generator._get_theme_color("LIGHT_1"), MSO_THEME_COLOR.LIGHT_1)
        self.assertEqual(self.generator._get_theme_color("LIGHT_2"), MSO_THEME_COLOR.LIGHT_2)
        self.assertEqual(self.generator._get_theme_color("DARK_1"), MSO_THEME_COLOR.DARK_1)
        self.assertEqual(self.generator._get_theme_color("DARK_2"), MSO_THEME_COLOR.DARK_2)
        self.assertEqual(self.generator._get_theme_color("ACCENT_1"), MSO_THEME_COLOR.ACCENT_1)
        self.assertEqual(self.generator._get_theme_color("ACCENT_2"), MSO_THEME_COLOR.ACCENT_2)
        self.assertEqual(self.generator._get_theme_color("ACCENT_3"), MSO_THEME_COLOR.ACCENT_3)
        self.assertEqual(self.generator._get_theme_color("ACCENT_4"), MSO_THEME_COLOR.ACCENT_4)
        self.assertEqual(self.generator._get_theme_color("ACCENT_5"), MSO_THEME_COLOR.ACCENT_5)
        self.assertEqual(self.generator._get_theme_color("ACCENT_6"), MSO_THEME_COLOR.ACCENT_6)

    def test_invalid_color_returns_default(self):
        """Invalid color name returns default LIGHT_2."""
        self.assertEqual(self.generator._get_theme_color("INVALID"), MSO_THEME_COLOR.LIGHT_2)
        self.assertEqual(self.generator._get_theme_color(""), MSO_THEME_COLOR.LIGHT_2)
        # Case sensitive - lowercase should return default
        self.assertEqual(self.generator._get_theme_color("light_1"), MSO_THEME_COLOR.LIGHT_2)

    def test_all_theme_colors_mapped(self):
        """All expected theme colors have mappings."""
        expected_colors = [
            "LIGHT_1", "LIGHT_2", "DARK_1", "DARK_2",
            "ACCENT_1", "ACCENT_2", "ACCENT_3", "ACCENT_4", "ACCENT_5", "ACCENT_6",
        ]
        for color_name in expected_colors:
            result = self.generator._get_theme_color(color_name)
            self.assertIsNotNone(result)
            self.assertIsInstance(result, MSO_THEME_COLOR)


class TestSlideGeneratorAddTextToParagraph(unittest.TestCase):
    """Tests for SlideGenerator._add_text_to_paragraph method."""

    def setUp(self):
        """Create a SlideGenerator instance for testing."""
        self.generator = SlideGenerator(template_path="/fake/template.pptx")

    def _create_mock_paragraph(self):
        """Create a mock paragraph with run tracking."""
        mock_para = MagicMock()
        mock_runs = []

        def add_run():
            mock_run = MagicMock()
            mock_run.text = ""
            mock_run.font = MagicMock()
            mock_run.font.size = None
            mock_run.font.color = MagicMock()
            mock_run.font.color.theme_color = None
            mock_run.font.bold = None
            mock_runs.append(mock_run)
            return mock_run

        mock_para.add_run = add_run
        mock_para._runs = mock_runs
        return mock_para, mock_runs

    def test_simple_text_no_highlights(self):
        """Simple text without highlights creates single run."""
        mock_para, runs = self._create_mock_paragraph()

        self.generator._add_text_to_paragraph(
            mock_para,
            "Simple text",
            TextStyle(
                theme_color=MSO_THEME_COLOR.LIGHT_2,
                font_size=Pt(15),
                bold=False,
                highlights=None,
            ),
        )

        self.assertEqual(len(runs), 1)
        self.assertEqual(runs[0].text, "Simple text")
        self.assertEqual(runs[0].font.color.theme_color, MSO_THEME_COLOR.LIGHT_2)
        self.assertIsNone(runs[0].font.bold)  # Inherits template default

    def test_text_with_single_highlight(self):
        """Text with single highlight creates multiple runs."""
        mock_para, runs = self._create_mock_paragraph()

        self.generator._add_text_to_paragraph(
            mock_para,
            "This has one highlight word here",
            TextStyle(
                theme_color=MSO_THEME_COLOR.LIGHT_2,
                font_size=Pt(15),
                bold=False,
                highlights=["highlight"],
            ),
        )

        # Should have 3 runs: "This has one " + "highlight" + " word here"
        self.assertEqual(len(runs), 3)

        # Regular text
        self.assertEqual(runs[0].text, "This has one ")
        self.assertEqual(runs[0].font.color.theme_color, MSO_THEME_COLOR.LIGHT_2)

        # Highlighted text
        self.assertEqual(runs[1].text, "highlight")
        self.assertEqual(runs[1].font.color.theme_color, HIGHLIGHT_THEME_COLOR)
        self.assertTrue(runs[1].font.bold)

        # Regular text
        self.assertEqual(runs[2].text, " word here")

    def test_text_with_multiple_highlights(self):
        """Text with multiple highlights creates correct runs."""
        mock_para, runs = self._create_mock_paragraph()

        self.generator._add_text_to_paragraph(
            mock_para,
            "First word and second word",
            TextStyle(
                theme_color=MSO_THEME_COLOR.LIGHT_2,
                font_size=Pt(15),
                bold=False,
                highlights=["First", "second"],
            ),
        )

        # Should have runs for: "First" + " word and " + "second" + " word"
        self.assertGreaterEqual(len(runs), 4)

        # Check highlighted runs are bold and accent colored
        highlighted_runs = [r for r in runs if r.text in ["First", "second"]]
        for run in highlighted_runs:
            self.assertTrue(run.font.bold)
            self.assertEqual(run.font.color.theme_color, HIGHLIGHT_THEME_COLOR)

    def test_header_is_bold(self):
        """Header text without highlights is bold."""
        mock_para, runs = self._create_mock_paragraph()

        self.generator._add_text_to_paragraph(
            mock_para,
            "Header Text:",
            TextStyle(
                theme_color=MSO_THEME_COLOR.LIGHT_2,
                font_size=Pt(18),
                bold=True,
                highlights=None,
            ),
        )

        self.assertEqual(len(runs), 1)
        self.assertTrue(runs[0].font.bold)

    def test_highlight_is_bold_and_accent(self):
        """Highlighted text is both bold and accent colored."""
        mock_para, runs = self._create_mock_paragraph()

        self.generator._add_text_to_paragraph(
            mock_para,
            "Check this important item",
            TextStyle(
                theme_color=MSO_THEME_COLOR.LIGHT_2,
                font_size=Pt(15),
                bold=False,
                highlights=["important"],
            ),
        )

        # Find the highlighted run
        important_run = next(r for r in runs if r.text == "important")
        self.assertTrue(important_run.font.bold)
        self.assertEqual(important_run.font.color.theme_color, HIGHLIGHT_THEME_COLOR)

    def test_empty_highlights_list(self):
        """Empty highlights list treated same as None."""
        mock_para, runs = self._create_mock_paragraph()

        self.generator._add_text_to_paragraph(
            mock_para,
            "Regular text",
            TextStyle(
                theme_color=MSO_THEME_COLOR.LIGHT_2,
                font_size=Pt(15),
                bold=False,
                highlights=[],
            ),
        )

        self.assertEqual(len(runs), 1)
        self.assertEqual(runs[0].text, "Regular text")

    def test_url_sets_hyperlink_on_simple_run(self):
        """When url is provided, the run gets hyperlink address, blue color, and underline."""
        mock_para, runs = self._create_mock_paragraph()

        self.generator._add_text_to_paragraph(
            mock_para,
            "Click here",
            TextStyle(
                theme_color=MSO_THEME_COLOR.LIGHT_2,
                font_size=Pt(15),
                bold=False,
                highlights=None,
                url="https://example.com",
            ),
        )

        self.assertEqual(len(runs), 1)
        self.assertEqual(runs[0].text, "Click here")
        self.assertEqual(runs[0].hyperlink.address, "https://example.com")
        self.assertTrue(runs[0].font.underline)
        self.assertEqual(runs[0].font.color.rgb, LINK_BLUE)

    def test_url_with_highlights_sets_hyperlink_on_all_runs(self):
        """When url is provided with highlights, all runs get hyperlink; highlighted runs keep their color."""
        mock_para, runs = self._create_mock_paragraph()

        self.generator._add_text_to_paragraph(
            mock_para,
            "Visit dashboard now",
            TextStyle(
                theme_color=MSO_THEME_COLOR.LIGHT_2,
                font_size=Pt(15),
                bold=False,
                highlights=["dashboard"],
                url="https://grafana.example.com",
            ),
        )

        # 3 runs: "Visit " + "dashboard" + " now"
        self.assertEqual(len(runs), 3)
        for run in runs:
            self.assertEqual(run.hyperlink.address, "https://grafana.example.com")
            self.assertTrue(run.font.underline)

    def test_no_url_does_not_set_hyperlink(self):
        """When url is None, hyperlink is not set on the run."""
        mock_para, runs = self._create_mock_paragraph()

        self.generator._add_text_to_paragraph(
            mock_para,
            "Plain text",
            TextStyle(
                theme_color=MSO_THEME_COLOR.LIGHT_2,
                font_size=Pt(15),
                bold=False,
                highlights=None,
                url=None,
            ),
        )

        self.assertEqual(len(runs), 1)
        self.assertEqual(runs[0].text, "Plain text")
        # _apply_hyperlink should not have been called — verify address was never set
        # (MagicMock auto-creates attributes, so we check it wasn't assigned a real URL)
        for run in runs:
            if hasattr(run.hyperlink.address, '_mock_name'):
                pass  # Mock attribute, never explicitly set — correct
            else:
                self.assertIsNone(run.hyperlink.address)


class TestSlideGeneratorFindShape(unittest.TestCase):
    """Tests for _find_shape helper method."""

    def test_find_placeholder_shape(self):
        """Test finding a placeholder shape."""
        generator = SlideGenerator(template_path="/template.pptx")

        mock_slide = MagicMock()
        mock_shape1 = MagicMock()
        mock_shape1.is_placeholder = False
        mock_shape2 = MagicMock()
        mock_shape2.is_placeholder = True
        mock_shape2.has_text_frame = True
        mock_slide.shapes = [mock_shape1, mock_shape2]

        result = generator._find_shape(mock_slide, placeholder=True, has_text_frame=True)
        self.assertEqual(result, mock_shape2)

    def test_find_non_placeholder_shape(self):
        """Test finding a non-placeholder shape."""
        generator = SlideGenerator(template_path="/template.pptx")

        mock_slide = MagicMock()
        mock_shape1 = MagicMock()
        mock_shape1.is_placeholder = True
        mock_shape2 = MagicMock()
        mock_shape2.is_placeholder = False
        mock_shape2.shape_type = MSO_SHAPE_TYPE.TEXT_BOX
        mock_slide.shapes = [mock_shape1, mock_shape2]

        result = generator._find_shape(
            mock_slide, placeholder=False, shape_type=MSO_SHAPE_TYPE.TEXT_BOX
        )
        self.assertEqual(result, mock_shape2)

    def test_find_shape_returns_none_when_not_found(self):
        """Test that _find_shape returns None when no match."""
        generator = SlideGenerator(template_path="/template.pptx")

        mock_slide = MagicMock()
        mock_shape = MagicMock()
        mock_shape.is_placeholder = False
        mock_shape.has_text_frame = False
        mock_slide.shapes = [mock_shape]

        result = generator._find_shape(mock_slide, placeholder=True, has_text_frame=True)
        self.assertIsNone(result)

    def test_find_shape_skips_wrong_shape_type(self):
        """Test that _find_shape skips shapes with wrong shape_type."""
        generator = SlideGenerator(template_path="/template.pptx")

        mock_slide = MagicMock()
        mock_shape = MagicMock()
        mock_shape.is_placeholder = False
        mock_shape.shape_type = MSO_SHAPE_TYPE.PICTURE  # Not TEXT_BOX
        mock_slide.shapes = [mock_shape]

        result = generator._find_shape(
            mock_slide, placeholder=False, shape_type=MSO_SHAPE_TYPE.TEXT_BOX
        )
        self.assertIsNone(result)

    def test_find_shape_skips_no_text_frame(self):
        """Test that _find_shape skips shapes without text frame when required."""
        generator = SlideGenerator(template_path="/template.pptx")

        mock_slide = MagicMock()
        mock_shape = MagicMock()
        mock_shape.is_placeholder = False
        mock_shape.shape_type = MSO_SHAPE_TYPE.TEXT_BOX
        mock_shape.has_text_frame = False  # No text frame
        mock_slide.shapes = [mock_shape]

        result = generator._find_shape(
            mock_slide,
            placeholder=False,
            shape_type=MSO_SHAPE_TYPE.TEXT_BOX,
            has_text_frame=True,
        )
        self.assertIsNone(result)


class TestSlideGeneratorStyleRun(unittest.TestCase):
    """Tests for _style_run helper method."""

    def test_style_run_sets_all_properties(self):
        """Test that _style_run sets font size, color, and bold."""
        generator = SlideGenerator(template_path="/template.pptx")
        mock_run = MagicMock()

        generator._style_run(
            mock_run,
            font_size=Pt(14),
            theme_color=MSO_THEME_COLOR.ACCENT_1,
            bold=True,
        )

        self.assertEqual(mock_run.font.size, Pt(14))
        self.assertEqual(mock_run.font.color.theme_color, MSO_THEME_COLOR.ACCENT_1)
        self.assertTrue(mock_run.font.bold)

    def test_style_run_skips_none_values(self):
        """Test that _style_run skips properties set to None."""
        generator = SlideGenerator(template_path="/template.pptx")
        mock_run = MagicMock()

        # Only set theme_color, leave others as None
        generator._style_run(mock_run, theme_color=MSO_THEME_COLOR.LIGHT_2)

        self.assertEqual(mock_run.font.color.theme_color, MSO_THEME_COLOR.LIGHT_2)
        # font.size and font.bold should not have been set
        # (MagicMock will create them on access, but we didn't explicitly set them)

    def test_style_run_handles_partial_styling(self):
        """Test styling with only font_size."""
        generator = SlideGenerator(template_path="/template.pptx")
        mock_run = MagicMock()

        generator._style_run(mock_run, font_size=Pt(18))

        self.assertEqual(mock_run.font.size, Pt(18))


class TestSlideGeneratorSetSlideTitle(unittest.TestCase):
    """Tests for _set_slide_title method."""

    def test_set_slide_title_applies_theme_color(self):
        """Test that _set_slide_title applies theme color to runs."""
        generator = SlideGenerator(template_path="/template.pptx")

        mock_slide = MagicMock()
        mock_shape = MagicMock()
        mock_shape.is_placeholder = True
        mock_shape.has_text_frame = True
        mock_paragraph = MagicMock()
        mock_run = MagicMock()
        mock_paragraph.runs = [mock_run]
        mock_shape.text_frame.paragraphs = [mock_paragraph]
        mock_slide.shapes = [mock_shape]

        generator._set_slide_title(mock_slide, "Test Title", MSO_THEME_COLOR.LIGHT_2)

        self.assertEqual(mock_paragraph.text, "Test Title")
        # Verify theme_color was set on the run's font color
        self.assertEqual(mock_run.font.color.theme_color, MSO_THEME_COLOR.LIGHT_2)


class TestSlideGeneratorRepositionTextbox(unittest.TestCase):
    """Tests for _reposition_textbox method."""

    def test_reposition_textbox_finds_and_moves_textbox(self):
        """Test that _reposition_textbox repositions non-placeholder text boxes."""
        generator = SlideGenerator(template_path="/template.pptx")

        mock_slide = MagicMock()
        mock_shape = MagicMock()
        mock_shape.is_placeholder = False
        mock_shape.shape_type = MSO_SHAPE_TYPE.TEXT_BOX
        mock_slide.shapes = [mock_shape]

        mock_shape.top = 914400 * 3.54  # 3.54" in EMU (template position)
        generator._reposition_textbox(mock_slide, 1.0, 7.0)

        # Verify position was set (left and width only; top preserved from template)
        self.assertEqual(mock_shape.left, Inches(1.0))
        self.assertEqual(mock_shape.width, Inches(7.0))

    def test_reposition_textbox_no_textbox_found(self):
        """Test that _reposition_textbox handles missing text box gracefully."""
        generator = SlideGenerator(template_path="/template.pptx")

        mock_slide = MagicMock()
        mock_shape = MagicMock()
        mock_shape.is_placeholder = True  # Only placeholder, no text box
        mock_slide.shapes = [mock_shape]

        # Should not raise an error
        generator._reposition_textbox(mock_slide, 1.0, 7.0)


class TestSlideGeneratorSetSlideContent(unittest.TestCase):
    """Tests for _set_slide_content method."""

    def test_set_slide_content_no_textbox_found(self):
        """Test _set_slide_content returns early when no text box shape found."""
        generator = SlideGenerator(template_path="/template.pptx")

        mock_slide = MagicMock()
        # Only placeholder shape, no text box
        mock_placeholder = MagicMock()
        mock_placeholder.is_placeholder = True
        mock_placeholder.has_text_frame = True
        mock_placeholder.text_frame = MagicMock()
        mock_placeholder.text_frame.paragraphs = [MagicMock(runs=[])]
        mock_slide.shapes = [mock_placeholder]

        # Should not raise, just return early
        generator._set_slide_content(
            mock_slide,
            "Title",
            ["Bullet 1"],
            MSO_THEME_COLOR.LIGHT_2,
        )

    def test_set_slide_content_with_bullet_items(self):
        """Test _set_slide_content handles BulletItem objects."""
        generator = SlideGenerator(template_path="/template.pptx")

        # Create mock slide with title and text box
        mock_slide = MagicMock()

        mock_title = MagicMock()
        mock_title.is_placeholder = True
        mock_title.has_text_frame = True
        mock_title.placeholder_format.idx = 0
        mock_title.text_frame = MagicMock()
        mock_title_para = MagicMock()
        mock_title_para.text = ""
        mock_title_para.runs = []
        mock_title.text_frame.paragraphs = [mock_title_para]

        mock_textbox = MagicMock()
        mock_textbox.is_placeholder = False
        mock_textbox.shape_type = MSO_SHAPE_TYPE.TEXT_BOX
        mock_textbox.has_text_frame = True

        mock_tf = MagicMock()
        mock_para = MagicMock()
        mock_tf.paragraphs = [mock_para]
        mock_tf.add_paragraph = MagicMock(return_value=MagicMock())
        mock_textbox.text_frame = mock_tf

        mock_slide.shapes = [mock_title, mock_textbox]

        # Use BulletItem objects
        bullets = [
            BulletItem(text="First bullet", level=0, highlight=["First"]),
            BulletItem(text="Second bullet", level=1),
        ]

        # Patch _apply_native_bullet since it requires real lxml elements
        with patch.object(generator, '_apply_native_bullet'):
            generator._set_slide_content(
                mock_slide,
                "Test Title",
                bullets,
                MSO_THEME_COLOR.LIGHT_2,
            )

        # Verify text frame was cleared
        mock_tf.clear.assert_called_once()

    def test_set_slide_content_with_section_header(self):
        """Test _set_slide_content applies header formatting for section headers."""
        generator = SlideGenerator(template_path="/template.pptx")

        mock_slide = MagicMock()

        mock_title = MagicMock()
        mock_title.is_placeholder = True
        mock_title.has_text_frame = True
        mock_title.text_frame = MagicMock()
        mock_title.text_frame.paragraphs = [MagicMock(runs=[])]

        mock_textbox = MagicMock()
        mock_textbox.is_placeholder = False
        mock_textbox.shape_type = MSO_SHAPE_TYPE.TEXT_BOX
        mock_textbox.has_text_frame = True

        mock_tf = MagicMock()
        mock_para = MagicMock()
        mock_tf.paragraphs = [mock_para]
        mock_textbox.text_frame = mock_tf

        mock_slide.shapes = [mock_title, mock_textbox]

        # Section header ends with colon and is short
        bullets = ["Overview:"]

        generator._set_slide_content(
            mock_slide,
            "Test Title",
            bullets,
            MSO_THEME_COLOR.LIGHT_2,
        )

        # Verify header spacing was applied
        self.assertEqual(mock_para.space_before, Pt(SPACING_BEFORE_HEADER))
        self.assertEqual(mock_para.space_after, Pt(SPACING_AFTER_HEADER))


class TestSlideGeneratorAddTableToSlide(unittest.TestCase):
    """Tests for _add_table_to_slide method."""

    def test_add_table_empty_headers(self):
        """Test _add_table_to_slide returns early with empty headers."""
        generator = SlideGenerator(template_path="/template.pptx")

        mock_slide = MagicMock()

        # Should return early and not add any table
        generator._add_table_to_slide(
            mock_slide,
            headers=[],  # Empty headers
            rows=[["A", "B"]],
            theme_color=MSO_THEME_COLOR.LIGHT_2,
        )

        mock_slide.shapes.add_table.assert_not_called()

    def test_add_table_empty_rows(self):
        """Test _add_table_to_slide returns early with empty rows."""
        generator = SlideGenerator(template_path="/template.pptx")

        mock_slide = MagicMock()

        # Should return early and not add any table
        generator._add_table_to_slide(
            mock_slide,
            headers=["Col1", "Col2"],
            rows=[],  # Empty rows
            theme_color=MSO_THEME_COLOR.LIGHT_2,
        )

        mock_slide.shapes.add_table.assert_not_called()

    def test_add_table_with_first_col_width(self):
        """Test _add_table_to_slide applies first column width when specified."""
        generator = SlideGenerator(template_path="/template.pptx")

        mock_slide = MagicMock()
        mock_table_shape = MagicMock()
        mock_table = MagicMock()
        mock_table_shape.table = mock_table
        mock_slide.shapes.add_table.return_value = mock_table_shape

        # Set up column mocks
        mock_col0 = MagicMock()
        mock_col1 = MagicMock()
        mock_col2 = MagicMock()
        mock_table.columns = [mock_col0, mock_col1, mock_col2]

        # Set up cell mocks
        mock_cell = MagicMock()
        mock_cell.text_frame = MagicMock()
        mock_cell.text_frame.paragraphs = [MagicMock(runs=[])]
        mock_table.cell = MagicMock(return_value=mock_cell)

        generator._add_table_to_slide(
            mock_slide,
            headers=["Name", "Value", "Status"],
            rows=[["A", "B", "C"]],
            theme_color=MSO_THEME_COLOR.LIGHT_2,
            first_col_width=2.5,  # Specify first column width
        )

        # Verify first column width was set
        self.assertEqual(mock_col0.width, int(Inches(2.5)))

    def test_add_table_styles_header_runs(self):
        """Test _add_table_to_slide styles header cell runs."""
        generator = SlideGenerator(template_path="/template.pptx")

        mock_slide = MagicMock()
        mock_table_shape = MagicMock()
        mock_table = MagicMock()
        mock_table_shape.table = mock_table
        mock_slide.shapes.add_table.return_value = mock_table_shape

        # Set up column mocks
        mock_table.columns = [MagicMock(), MagicMock()]

        # Set up cell with runs
        mock_run = MagicMock()
        mock_para = MagicMock()
        mock_para.runs = [mock_run]
        mock_cell = MagicMock()
        mock_cell.text_frame = MagicMock()
        mock_cell.text_frame.paragraphs = [mock_para]
        mock_table.cell = MagicMock(return_value=mock_cell)

        generator._add_table_to_slide(
            mock_slide,
            headers=["Col1", "Col2"],
            rows=[["A", "B"]],
            theme_color=MSO_THEME_COLOR.LIGHT_2,
        )

        # Verify _style_run was called on header runs (bold=True)
        self.assertTrue(mock_run.font.bold)


class TestNormalizeTableRows(unittest.TestCase):
    """Tests for SlideGenerator._normalize_table_rows method."""

    def test_empty_rows_skipped(self):
        """Empty rows (None-ish / []) are skipped."""
        generator = SlideGenerator(template_path="/template.pptx")
        rows = [[], ["a", "b"], [], ["c", "d"]]
        result = generator._normalize_table_rows(rows, num_cols=2)
        self.assertEqual(result, [["a", "b"], ["c", "d"]])

    def test_short_rows_padded(self):
        """Rows shorter than num_cols are padded with empty strings."""
        generator = SlideGenerator(template_path="/template.pptx")
        rows = [["a"], ["b", "c"]]
        result = generator._normalize_table_rows(rows, num_cols=3)
        self.assertEqual(result, [["a", "", ""], ["b", "c", ""]])

    def test_long_rows_truncated(self):
        """Rows longer than num_cols are truncated."""
        generator = SlideGenerator(template_path="/template.pptx")
        rows = [["a", "b", "c", "d"]]
        result = generator._normalize_table_rows(rows, num_cols=2)
        self.assertEqual(result, [["a", "b"]])

    def test_all_empty_returns_empty(self):
        """All empty rows returns an empty list."""
        generator = SlideGenerator(template_path="/template.pptx")
        rows = [[], [], []]
        result = generator._normalize_table_rows(rows, num_cols=3)
        self.assertEqual(result, [])

    def test_no_rows_returns_empty(self):
        """An empty input list returns an empty list."""
        generator = SlideGenerator(template_path="/template.pptx")
        result = generator._normalize_table_rows([], num_cols=2)
        self.assertEqual(result, [])

    def test_exact_length_rows_unchanged(self):
        """Rows with exactly num_cols columns pass through unchanged."""
        generator = SlideGenerator(template_path="/template.pptx")
        rows = [["a", "b", "c"], ["d", "e", "f"]]
        result = generator._normalize_table_rows(rows, num_cols=3)
        self.assertEqual(result, [["a", "b", "c"], ["d", "e", "f"]])

    def test_mixed_lengths(self):
        """Mix of short, exact, and long rows are all normalized."""
        generator = SlideGenerator(template_path="/template.pptx")
        rows = [["short"], ["a", "b"], ["a", "b", "c", "extra"]]
        result = generator._normalize_table_rows(rows, num_cols=2)
        self.assertEqual(result, [["short", ""], ["a", "b"], ["a", "b"]])


class TestSetTableColumnWidths(unittest.TestCase):
    """Tests for SlideGenerator._set_table_column_widths method."""

    def test_equal_widths_without_first_col_width(self):
        """Without first_col_width, all columns get equal width."""
        generator = SlideGenerator(template_path="/template.pptx")
        num_cols = 3
        total_width = 9000000  # EMU

        mock_columns = [MagicMock() for _ in range(num_cols)]
        mock_table = MagicMock()
        mock_table.columns.__getitem__ = lambda self, i: mock_columns[i]

        generator._set_table_column_widths(mock_table, num_cols, total_width, first_col_width=None)

        expected_width = int(total_width / num_cols)
        for col in mock_columns:
            self.assertEqual(col.width, expected_width)

    def test_first_col_width_set(self):
        """With first_col_width, first column gets specified width, rest share remainder."""
        generator = SlideGenerator(template_path="/template.pptx")
        num_cols = 3
        total_width = 9144000  # EMU (approximately 10 inches)
        first_col_inches = 2.0

        mock_columns = [MagicMock() for _ in range(num_cols)]
        mock_table = MagicMock()
        mock_table.columns.__getitem__ = lambda self, i: mock_columns[i]

        generator._set_table_column_widths(
            mock_table, num_cols, total_width, first_col_width=first_col_inches
        )

        first_width = Inches(first_col_inches)
        other_width = int((total_width - first_width) / (num_cols - 1))
        self.assertEqual(mock_columns[0].width, int(first_width))
        for col in mock_columns[1:]:
            self.assertEqual(col.width, other_width)

    def test_single_column_with_first_col_width(self):
        """With num_cols=1, first_col_width is ignored (falls to equal-width branch)."""
        generator = SlideGenerator(template_path="/template.pptx")
        num_cols = 1
        total_width = 9000000

        mock_columns = [MagicMock()]
        mock_table = MagicMock()
        mock_table.columns.__getitem__ = lambda self, i: mock_columns[i]

        generator._set_table_column_widths(
            mock_table, num_cols, total_width, first_col_width=3.0
        )

        # num_cols <= 1 so falls to else branch: equal widths
        self.assertEqual(mock_columns[0].width, int(total_width / num_cols))

    def test_equal_widths_when_first_col_zero(self):
        """first_col_width=0 means no override, so equal widths are used."""
        generator = SlideGenerator(template_path="/template.pptx")
        num_cols = 2
        total_width = 8000000

        mock_columns = [MagicMock() for _ in range(num_cols)]
        mock_table = MagicMock()
        mock_table.columns.__getitem__ = lambda self, i: mock_columns[i]

        generator._set_table_column_widths(
            mock_table, num_cols, total_width, first_col_width=0
        )

        expected_width = int(total_width / num_cols)
        for col in mock_columns:
            self.assertEqual(col.width, expected_width)


class TestStyleTableHeader(unittest.TestCase):
    """Tests for SlideGenerator._style_table_header method."""

    def _make_cell(self, num_runs=1):
        """Create a mock table cell with text_frame, paragraphs, and runs."""
        cell = MagicMock()
        runs = [MagicMock() for _ in range(num_runs)]
        paragraph = MagicMock()
        paragraph.runs = runs
        cell.text_frame.paragraphs = [paragraph]
        return cell

    def test_sets_header_text_and_fill(self):
        """Verify header text, background fill, and vertical anchor are set."""
        generator = SlideGenerator(template_path="/template.pptx")
        headers = ["Name", "Value", "Status"]
        cells = [self._make_cell() for _ in headers]
        mock_table = MagicMock()
        mock_table.cell = lambda row, col: cells[col]

        theme_color = MagicMock()

        with patch.object(generator, '_style_run'):
            generator._style_table_header(mock_table, headers, theme_color)

        for j, header in enumerate(headers):
            cell = cells[j]
            self.assertEqual(cell.text, header)
            cell.fill.solid.assert_called_once()
            self.assertEqual(cell.fill.fore_color.rgb, TABLE_HEADER_BG)
            self.assertEqual(cell.vertical_anchor, VERTICAL_ANCHOR_MIDDLE)

    def test_first_column_left_aligned(self):
        """First column paragraph is LEFT aligned."""
        from pptx.enum.text import PP_ALIGN

        generator = SlideGenerator(template_path="/template.pptx")
        headers = ["Name", "Value"]
        cells = [self._make_cell() for _ in headers]
        mock_table = MagicMock()
        mock_table.cell = lambda row, col: cells[col]

        with patch.object(generator, '_style_run'):
            generator._style_table_header(mock_table, headers, MagicMock())

        first_paragraph = cells[0].text_frame.paragraphs[0]
        self.assertEqual(first_paragraph.alignment, PP_ALIGN.LEFT)

    def test_non_first_columns_center_aligned(self):
        """Non-first column paragraphs are CENTER aligned."""
        from pptx.enum.text import PP_ALIGN

        generator = SlideGenerator(template_path="/template.pptx")
        headers = ["Name", "Value", "Status"]
        cells = [self._make_cell() for _ in headers]
        mock_table = MagicMock()
        mock_table.cell = lambda row, col: cells[col]

        with patch.object(generator, '_style_run'):
            generator._style_table_header(mock_table, headers, MagicMock())

        for j in range(1, len(headers)):
            paragraph = cells[j].text_frame.paragraphs[0]
            self.assertEqual(paragraph.alignment, PP_ALIGN.CENTER)

    def test_style_run_called_with_bold(self):
        """_style_run is called with bold=True for each header run."""
        generator = SlideGenerator(template_path="/template.pptx")
        headers = ["Col1"]
        cells = [self._make_cell(num_runs=1)]
        mock_table = MagicMock()
        mock_table.cell = lambda row, col: cells[col]
        theme_color = MagicMock()

        with patch.object(generator, '_style_run') as mock_style_run:
            generator._style_table_header(mock_table, headers, theme_color)

        self.assertTrue(mock_style_run.called)
        call_kwargs = mock_style_run.call_args
        self.assertTrue(call_kwargs[1].get('bold') or call_kwargs.kwargs.get('bold'))


class TestStyleTableDataRows(unittest.TestCase):
    """Tests for SlideGenerator._style_table_data_rows method."""

    def _make_cell(self):
        """Create a mock table cell with text_frame and a paragraph with add_run."""
        cell = MagicMock()
        mock_run = MagicMock()
        paragraph = MagicMock()
        paragraph.add_run.return_value = mock_run
        cell.text_frame.paragraphs = [paragraph]
        return cell, mock_run

    def test_alternating_row_backgrounds(self):
        """Even rows get TABLE_ROW_EVEN_BG, odd rows get TABLE_ROW_ODD_BG."""
        generator = SlideGenerator(template_path="/template.pptx")
        rows = [["a", "b"], ["c", "d"], ["e", "f"]]

        cells = {}
        for i in range(len(rows)):
            for j in range(2):
                cell, _ = self._make_cell()
                cells[(i + 1, j)] = cell

        mock_table = MagicMock()
        mock_table.cell = lambda row, col: cells[(row, col)]

        with patch.object(generator, '_style_run'):
            generator._style_table_data_rows(mock_table, rows, MagicMock())

        # Row 0 (even): TABLE_ROW_EVEN_BG
        self.assertEqual(cells[(1, 0)].fill.fore_color.rgb, TABLE_ROW_EVEN_BG)
        self.assertEqual(cells[(1, 1)].fill.fore_color.rgb, TABLE_ROW_EVEN_BG)
        # Row 1 (odd): TABLE_ROW_ODD_BG
        self.assertEqual(cells[(2, 0)].fill.fore_color.rgb, TABLE_ROW_ODD_BG)
        self.assertEqual(cells[(2, 1)].fill.fore_color.rgb, TABLE_ROW_ODD_BG)
        # Row 2 (even): TABLE_ROW_EVEN_BG
        self.assertEqual(cells[(3, 0)].fill.fore_color.rgb, TABLE_ROW_EVEN_BG)
        self.assertEqual(cells[(3, 1)].fill.fore_color.rgb, TABLE_ROW_EVEN_BG)

    def test_vertical_anchor_set(self):
        """All data cells get VERTICAL_ANCHOR_MIDDLE."""
        generator = SlideGenerator(template_path="/template.pptx")
        rows = [["a"]]

        cell, _ = self._make_cell()
        mock_table = MagicMock()
        mock_table.cell = lambda row, col: cell

        with patch.object(generator, '_style_run'):
            generator._style_table_data_rows(mock_table, rows, MagicMock())

        self.assertEqual(cell.vertical_anchor, VERTICAL_ANCHOR_MIDDLE)

    def test_first_column_left_aligned(self):
        """First column (j=0) is LEFT aligned."""
        from pptx.enum.text import PP_ALIGN

        generator = SlideGenerator(template_path="/template.pptx")
        rows = [["a", "b"]]

        cells = {}
        for j in range(2):
            cell, _ = self._make_cell()
            cells[(1, j)] = cell

        mock_table = MagicMock()
        mock_table.cell = lambda row, col: cells[(row, col)]

        with patch.object(generator, '_style_run'):
            generator._style_table_data_rows(mock_table, rows, MagicMock())

        p0 = cells[(1, 0)].text_frame.paragraphs[0]
        self.assertEqual(p0.alignment, PP_ALIGN.LEFT)

    def test_non_first_column_center_aligned(self):
        """Non-first columns (j>0) are CENTER aligned."""
        from pptx.enum.text import PP_ALIGN

        generator = SlideGenerator(template_path="/template.pptx")
        rows = [["a", "b", "c"]]

        cells = {}
        for j in range(3):
            cell, _ = self._make_cell()
            cells[(1, j)] = cell

        mock_table = MagicMock()
        mock_table.cell = lambda row, col: cells[(row, col)]

        with patch.object(generator, '_style_run'):
            generator._style_table_data_rows(mock_table, rows, MagicMock())

        for j in range(1, 3):
            p = cells[(1, j)].text_frame.paragraphs[0]
            self.assertEqual(p.alignment, PP_ALIGN.CENTER)

    def test_run_text_set_to_string_value(self):
        """Run text is set to str(value) for each cell."""
        generator = SlideGenerator(template_path="/template.pptx")
        rows = [["hello", 42]]

        cells = {}
        runs = {}
        for j in range(2):
            cell, run = self._make_cell()
            cells[(1, j)] = cell
            runs[(1, j)] = run

        mock_table = MagicMock()
        mock_table.cell = lambda row, col: cells[(row, col)]

        with patch.object(generator, '_style_run'):
            generator._style_table_data_rows(mock_table, rows, MagicMock())

        self.assertEqual(runs[(1, 0)].text, "hello")
        self.assertEqual(runs[(1, 1)].text, "42")

    def test_style_run_called_for_each_cell(self):
        """_style_run is called once per data cell."""
        generator = SlideGenerator(template_path="/template.pptx")
        rows = [["a", "b"], ["c", "d"]]

        cells = {}
        for i in range(2):
            for j in range(2):
                cell, _ = self._make_cell()
                cells[(i + 1, j)] = cell

        mock_table = MagicMock()
        mock_table.cell = lambda row, col: cells[(row, col)]

        with patch.object(generator, '_style_run') as mock_style_run:
            generator._style_table_data_rows(mock_table, rows, MagicMock())

        # 2 rows * 2 cols = 4 calls
        self.assertEqual(mock_style_run.call_count, 4)

    def test_empty_rows_no_cells_styled(self):
        """Empty rows list results in no cells being styled."""
        generator = SlideGenerator(template_path="/template.pptx")
        mock_table = MagicMock()

        with patch.object(generator, '_style_run') as mock_style_run:
            generator._style_table_data_rows(mock_table, [], MagicMock())

        mock_style_run.assert_not_called()


class TestAddBulletsBelow(unittest.TestCase):
    """Cover _add_bullets_below method fully (lines 564-591)."""

    def setUp(self) -> None:
        self.generator = SlideGenerator(template_path="/fake/template.pptx")

    def _make_mock_slide(self) -> tuple:
        """Return (mock_slide, mock_text_frame) with paragraph tracking."""
        mock_slide = MagicMock()
        mock_tf = MagicMock()
        mock_tf.paragraphs = [MagicMock()]

        mock_textbox = MagicMock()
        mock_textbox.text_frame = mock_tf
        mock_slide.shapes.add_textbox = MagicMock(return_value=mock_textbox)

        return mock_slide, mock_tf

    def test_adds_string_bullets(self) -> None:
        """String bullets are formatted as level-0 bullets."""
        mock_slide, mock_tf = self._make_mock_slide()

        self.generator._add_bullets_below(
            mock_slide,
            ["First note", "Second note"],
            MSO_THEME_COLOR.LIGHT_2,
            top_inches=3.0,
        )

        # Textbox should be added to slide
        mock_slide.shapes.add_textbox.assert_called_once()
        # Should have word_wrap enabled
        self.assertTrue(mock_tf.word_wrap)

    def test_adds_bullet_item_objects(self) -> None:
        """BulletItem objects are formatted with their level and highlights."""
        mock_slide, _ = self._make_mock_slide()

        bullets = [
            BulletItem(text="Highlighted item", level=1, highlight=["Highlighted"]),
        ]

        with patch.object(self.generator, "_add_text_to_paragraph") as mock_add_text:
            self.generator._add_bullets_below(
                mock_slide,
                bullets,
                MSO_THEME_COLOR.LIGHT_2,
                top_inches=2.5,
            )
            mock_add_text.assert_called_once()
            style = mock_add_text.call_args.args[2]
            # Check highlights were passed through
            self.assertFalse(style.bold)
            self.assertEqual(style.highlights, ["Highlighted"])

    def test_section_header_bullet_gets_header_font_size(self) -> None:
        """A bullet ending with ':' is treated as a section header."""
        mock_slide, _ = self._make_mock_slide()

        bullets = [BulletItem(text="Details:", level=0)]

        with patch.object(self.generator, "_add_text_to_paragraph") as mock_add_text:
            self.generator._add_bullets_below(
                mock_slide,
                bullets,
                MSO_THEME_COLOR.LIGHT_2,
                top_inches=2.0,
            )
            mock_add_text.assert_called_once()
            style = mock_add_text.call_args.args[2]
            # Font size should be header size
            self.assertEqual(style.font_size, Pt(FONT_SIZE_HEADER))
            # bold should be True
            self.assertTrue(style.bold)

    def test_mixed_string_and_bullet_items(self) -> None:
        """Mix of strings and BulletItem objects processes correctly."""
        mock_slide, _ = self._make_mock_slide()

        bullets = [
            "Plain string",
            BulletItem(text="Level 2 item", level=2),
        ]

        with patch.object(self.generator, "_add_text_to_paragraph") as mock_add_text:
            self.generator._add_bullets_below(
                mock_slide,
                bullets,
                MSO_THEME_COLOR.LIGHT_2,
                top_inches=3.5,
            )
            self.assertEqual(mock_add_text.call_count, 2)

    def test_remaining_height_minimum(self) -> None:
        """When top_inches is very high, remaining height is clamped to 0.5."""
        mock_slide, _ = self._make_mock_slide()

        self.generator._add_bullets_below(
            mock_slide,
            ["A note"],
            MSO_THEME_COLOR.LIGHT_2,
            top_inches=6.0,  # > 5.0, so remaining = max(5.0 - 6.0, 0.5) = 0.5
        )

        mock_slide.shapes.add_textbox.assert_called_once()

    def test_spacing_set_on_paragraphs(self) -> None:
        """Space before/after is set for each bullet paragraph."""
        mock_slide, mock_tf = self._make_mock_slide()
        mock_para = mock_tf.paragraphs[0]

        self.generator._add_bullets_below(
            mock_slide,
            ["Single bullet"],
            MSO_THEME_COLOR.LIGHT_2,
            top_inches=2.0,
        )

        self.assertEqual(mock_para.space_before, Pt(SPACING_BEFORE_BULLET))
        self.assertEqual(mock_para.space_after, Pt(SPACING_AFTER_BULLET))


class TestApplyNativeBullet(unittest.TestCase):
    """Cover _apply_native_bullet method (lines 369-396)."""

    def setUp(self) -> None:
        self.generator = SlideGenerator(template_path="/fake/template.pptx")
        self.nsmap = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        self.ns = self.nsmap["a"]

    def test_creates_ppr_and_applies_level0_bullet(self) -> None:
        """Creates pPr when missing and applies level-0 bullet character."""
        p_elem = etree.Element(f"{{{self.ns}}}p")
        etree.SubElement(p_elem, f"{{{self.ns}}}r")

        paragraph = MagicMock()
        paragraph._p = p_elem

        self.generator._apply_native_bullet(paragraph, level=0)

        p_pr = p_elem.find("a:pPr", self.nsmap)
        self.assertIsNotNone(p_pr)
        # pPr should be inserted as first child
        self.assertEqual(list(p_elem)[0].tag, f"{{{self.ns}}}pPr")
        # Level attribute
        self.assertEqual(p_pr.get("lvl"), "0")
        # Indent and margin
        self.assertEqual(p_pr.get("indent"), str(-457200))
        self.assertEqual(p_pr.get("marL"), str(457200))
        # Bullet character
        bu_char = p_pr.find("a:buChar", self.nsmap)
        self.assertIsNotNone(bu_char)
        self.assertEqual(bu_char.get("char"), "•")  # •

    def test_applies_level1_bullet(self) -> None:
        """Level-1 bullet uses open circle and correct margin."""
        p_elem = etree.Element(f"{{{self.ns}}}p")
        paragraph = MagicMock()
        paragraph._p = p_elem

        self.generator._apply_native_bullet(paragraph, level=1)

        p_pr = p_elem.find("a:pPr", self.nsmap)
        self.assertEqual(p_pr.get("lvl"), "1")
        self.assertEqual(p_pr.get("marL"), str(914400))
        bu_char = p_pr.find("a:buChar", self.nsmap)
        self.assertEqual(bu_char.get("char"), "◦")  # ◦

    def test_applies_level2_bullet(self) -> None:
        """Level-2 bullet uses small square and correct margin."""
        p_elem = etree.Element(f"{{{self.ns}}}p")
        paragraph = MagicMock()
        paragraph._p = p_elem

        self.generator._apply_native_bullet(paragraph, level=2)

        p_pr = p_elem.find("a:pPr", self.nsmap)
        self.assertEqual(p_pr.get("lvl"), "2")
        self.assertEqual(p_pr.get("marL"), str(1371600))
        bu_char = p_pr.find("a:buChar", self.nsmap)
        self.assertEqual(bu_char.get("char"), "▪")  # ▪

    def test_uses_existing_ppr(self) -> None:
        """When pPr already exists, reuses it instead of creating new one."""
        p_elem = etree.Element(f"{{{self.ns}}}p")
        existing_ppr = etree.SubElement(p_elem, f"{{{self.ns}}}pPr")

        paragraph = MagicMock()
        paragraph._p = p_elem

        self.generator._apply_native_bullet(paragraph, level=0)

        # Should still be only one pPr
        pprs = p_elem.findall("a:pPr", self.nsmap)
        self.assertEqual(len(pprs), 1)
        self.assertIs(pprs[0], existing_ppr)

    def test_removes_existing_bunone_before_adding_buchar(self) -> None:
        """Existing buNone is removed before adding buChar."""
        p_elem = etree.Element(f"{{{self.ns}}}p")
        p_pr = etree.SubElement(p_elem, f"{{{self.ns}}}pPr")
        etree.SubElement(p_pr, f"{{{self.ns}}}buNone")

        paragraph = MagicMock()
        paragraph._p = p_elem

        self.generator._apply_native_bullet(paragraph, level=0)

        bu_none = p_pr.find("a:buNone", self.nsmap)
        self.assertIsNone(bu_none)
        bu_char = p_pr.find("a:buChar", self.nsmap)
        self.assertIsNotNone(bu_char)

    def test_replaces_existing_buchar(self) -> None:
        """Existing buChar is removed and replaced with new one."""
        p_elem = etree.Element(f"{{{self.ns}}}p")
        p_pr = etree.SubElement(p_elem, f"{{{self.ns}}}pPr")
        old_bu = etree.SubElement(p_pr, f"{{{self.ns}}}buChar")
        old_bu.set("char", "X")

        paragraph = MagicMock()
        paragraph._p = p_elem

        self.generator._apply_native_bullet(paragraph, level=1)

        bu_chars = p_pr.findall("a:buChar", self.nsmap)
        self.assertEqual(len(bu_chars), 1)
        self.assertEqual(bu_chars[0].get("char"), "◦")

    def test_fallback_bullet_for_unknown_level(self) -> None:
        """Unknown level falls back to default bullet and indent."""
        p_elem = etree.Element(f"{{{self.ns}}}p")
        paragraph = MagicMock()
        paragraph._p = p_elem

        self.generator._apply_native_bullet(paragraph, level=5)

        p_pr = p_elem.find("a:pPr", self.nsmap)
        self.assertEqual(p_pr.get("lvl"), "5")
        # Falls back to default 457200
        self.assertEqual(p_pr.get("indent"), str(-457200))
        self.assertEqual(p_pr.get("marL"), str(457200))
        bu_char = p_pr.find("a:buChar", self.nsmap)
        self.assertEqual(bu_char.get("char"), "•")  # default •


class TestSetSlideTitleInheritStyle(unittest.TestCase):
    """Cover _set_slide_title with inherit_style=True (lines 271-273, 276-317)."""

    def setUp(self) -> None:
        self.generator = SlideGenerator(template_path="/fake/template.pptx")

    def _make_title_shape(self) -> MagicMock:
        """Return a mock title placeholder with paragraphs and runs."""
        shape = MagicMock()
        shape.is_placeholder = True
        shape.has_text_frame = True
        shape.placeholder_format.idx = 0
        shape.top = "SENTINEL_TOP"
        shape.left = "SENTINEL_LEFT"
        shape.width = "SENTINEL_WIDTH"
        shape.height = "SENTINEL_HEIGHT"

        para = MagicMock()
        para.text = ""
        para.runs = [MagicMock()]
        shape.text_frame.paragraphs = [para]
        return shape

    def test_inherit_style_skips_alignment_and_positioning(self) -> None:
        """inherit_style=True skips alignment override and does not reposition."""
        title_shape = self._make_title_shape()
        shapes = [title_shape]
        mock_slide = MagicMock()
        mock_slide.shapes.__iter__ = MagicMock(side_effect=lambda: iter(shapes))

        self.generator._set_slide_title(
            mock_slide, "Title", MSO_THEME_COLOR.LIGHT_2, inherit_style=True
        )

        # Title text is set
        self.assertEqual(title_shape.text_frame.paragraphs[0].text, "Title")
        # Positioning sentinel values are NOT overwritten
        self.assertEqual(title_shape.top, "SENTINEL_TOP")
        self.assertEqual(title_shape.left, "SENTINEL_LEFT")

    def test_inherit_style_false_does_reposition(self) -> None:
        """inherit_style=False (default) overwrites positioning."""
        title_shape = self._make_title_shape()
        shapes = [title_shape]
        mock_slide = MagicMock()
        mock_slide.shapes.__iter__ = MagicMock(side_effect=lambda: iter(shapes))

        self.generator._set_slide_title(
            mock_slide, "Title", MSO_THEME_COLOR.LIGHT_2, inherit_style=False
        )

        # Positioning was overwritten
        self.assertNotEqual(title_shape.top, "SENTINEL_TOP")


if __name__ == "__main__":
    unittest.main()
